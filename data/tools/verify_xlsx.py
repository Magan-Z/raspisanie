"""Независимая сверка: читаем xlsx напрямую (кодом, не связанным с парсером) и сравниваем с JSON-файлами групп.

Запуск:  python tools/verify_xlsx.py sources/raspisanie.xlsx public
Находит: потерянные и лишние занятия, неверные день / пару / неделю / подгруппу.
Тот же код вызывается тестом tests/test_crosscheck.py.
"""
import collections
import glob
import json
import re
import sys

import openpyxl

SHEETS = {"ОФО": "ofo", "ОЗФО": "ozfo", "Магистратура ОФО": "mag"}
DAYS = {"понедельник": 1, "вторник": 2, "среда": 3, "четверг": 4, "пятница": 5, "суббота": 6}


def norm(v):
    return " ".join(str(v).split()) if v is not None else ""


def check(xlsx: str, out: str) -> tuple[dict, list[str]]:
    """Возвращает (статистика, список расхождений). Пустой список — всё сходится."""
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    docs = {}
    for f in glob.glob(f"{out}/groups/*.json"):
        d = json.load(open(f, encoding="utf-8"))
        docs[(d["group"]["form"], d["group"]["title"])] = d
    problems, checked_cells, checked_groups = [], 0, 0
    for sheet, form in SHEETS.items():
        ws = wb[sheet]
        # своя карта объединений: (row, col) -> (r1, c1, r2, c2)
        area = {}
        for rng in ws.merged_cells.ranges:
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    area[(r, c)] = (rng.min_row, rng.min_col, rng.max_row, rng.max_col)
        def A(r, c): return area.get((r, c), (r, c, r, c))
        def text(r, c):
            r0, c0, _, _ = A(r, c)
            return norm(ws.cell(r0, c0).value)

        SKIP = re.compile(r"^(Дни недели|Часы)$|\d+\s*КУРС", re.I)
        # заголовки: строки, где встречается «Дни недели»; у каждой таблицы своя колонка дней
        header_cells = [(r, c) for r in range(1, ws.max_row + 1) for c in range(1, ws.max_column + 1) if text(r, c) == "Дни недели" and A(r, c)[0] == r and A(r, c)[1] == c]
        day_cols = sorted({c for _, c in header_cells})
        hrows = sorted({r for r, _ in header_cells})
        for dc in day_cols:
            rows_here = [r for r, c in header_cells if c == dc]
            for hi, h in enumerate(rows_here):
                end = (rows_here[hi + 1] - 1) if hi + 1 < len(rows_here) else ws.max_row
                # границы таблицы по столбцам: от dc до следующей колонки дней
                nxt = [c for c in day_cols if c > dc]
                c_end = (nxt[0] - 1) if nxt else ws.max_column
                def find_groups(row):
                    seen, found = set(), []
                    for c in range(dc + 2, c_end + 1):
                        t = text(row, c); a = A(row, c)
                        if t and not SKIP.search(t) and a not in seen:
                            seen.add(a); found.append((re.sub(r"\s*\([^)]*\)\s*$", "", t).strip(), a[1], a[3]))
                    return found
                groups = find_groups(h)
                grow = h
                if not groups:
                    groups = find_groups(h + 1); grow = h + 1
                slots, day = [], None
                first = grow + 1
                has_sub_row = not re.fullmatch(r"\d+-\d+", text(grow + 1, dc + 1) or "x")
                if has_sub_row: first = grow + 2
                for r in range(first, end + 1):
                    dt = text(r, dc).lower()
                    if dt in DAYS: day = DAYS[dt]
                    m = re.fullmatch(r"(\d+)-(\d+)", text(r, dc + 1))
                    if m and day and A(r, dc + 1)[0] == r:
                        a2 = A(r, dc + 1)
                        slots.append((day, (int(m.group(1)) + 1) // 2, a2[0], a2[2]))
                for sl in slots:
                    if sl[3] - sl[2] != 1: problems.append(f"{sheet}: слот {sl} не из двух строк")
                for name, c1, c2 in groups:
                    doc = docs.get((form, name))
                    if doc is None:
                        problems.append(f"{sheet}: группа «{name}» есть в Excel, но нет в JSON"); continue
                    checked_groups += 1
                    ncols = c2 - c1 + 1
                    cols = list(range(c1, c2 + 1))
                    labelled = has_sub_row and any(text(grow + 1, c) and A(grow + 1, c)[0] == grow + 1 for c in cols)
                    expected = collections.defaultdict(set)
                    for (day, pair, top, bottom) in slots:
                        for n, c in enumerate(cols, start=1):
                            t1, t2 = text(top, c), text(bottom, c)
                            if t1 and t1 == t2: expected[(day, pair, 0, t1)].add(n)
                            else:
                                if t1: expected[(day, pair, 1, t1)].add(n)
                                if t2: expected[(day, pair, 2, t2)].add(n)
                    got = collections.defaultdict(set)
                    nsub = len(doc["group"]["subgroups"])
                    for l in doc["lessons"]:
                        got[(l["weekday"], l["pair"], l["week"], l["raw"])].update(l["subgroups"])
                    if nsub == 1 and ncols > 1:
                        if labelled: problems.append(f"{name}: в Excel есть подписи подгрупп, а в JSON одна подгруппа")
                        for k, v in expected.items():
                            if v != set(range(1, ncols + 1)):
                                problems.append(f"{name}: подгруппы свёрнуты в одну, но {k} только у столбцов {sorted(v)}")
                        expected = {k: {1} for k in expected}
                    elif nsub != ncols:
                        problems.append(f"{name}: подгрупп в JSON {nsub}, столбцов в Excel {ncols}")
                    checked_cells += sum(len(v) for v in expected.values())
                    for k in sorted(set(expected) | set(got)):
                        if expected.get(k) != got.get(k):
                            problems.append(f"{sheet} / {name}: {k} -> Excel: {sorted(expected.get(k, []))}, JSON: {sorted(got.get(k, []))}")
    stats = {"groups": checked_groups, "json_groups": len(docs), "cells": checked_cells}
    return stats, problems


if __name__ == "__main__":
    stats, problems = check(sys.argv[1], sys.argv[2])
    print(f"Групп сверено: {stats['groups']} (в JSON групп: {stats['json_groups']})")
    print(f"Занятий (ячейка × подгруппа) сверено: {stats['cells']}")
    print(f"Расхождений: {len(problems)}")
    for p in problems[:60]:
        print(" -", p)
    sys.exit(1 if problems else 0)
