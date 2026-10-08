"""Проверки перед публикацией и сравнение версий."""

import copy
from pathlib import Path

import openpyxl
import pytest

from schedule_parser.diff import compare
from schedule_parser.extract import extract_workbook
from schedule_parser.layout import LayoutError
from schedule_parser.models import Group
from schedule_parser.publish import group_documents
from schedule_parser.validate import validate

DATA_DIR = Path(__file__).resolve().parent.parent


def previous_from(sheets: dict[str, list[Group]]) -> dict:
    docs = group_documents(sheets, "old")
    return {"index": {"version": "old"}, "groups": docs}


def test_real_file_is_valid(sheets: dict[str, list[Group]]) -> None:
    assert validate(sheets, None) == []
    assert validate(sheets, previous_from(sheets)) == []


def test_broken_file_without_group_row_fails(tmp_path, config: dict) -> None:
    # «Сломанный» файл: стираем строку с названиями групп на листе ОФО
    workbook = openpyxl.load_workbook(DATA_DIR / config["source"])
    sheet = workbook["ОФО"]
    for rng in list(sheet.merged_cells.ranges):
        if rng.min_row <= 7 and rng.max_row >= 6:
            sheet.unmerge_cells(str(rng))
    for row in (6, 7):
        for col in range(1, sheet.max_column + 1):
            if sheet.cell(row, col).value and "дни недели" not in str(sheet.cell(row, col).value).lower():
                sheet.cell(row, col).value = None
    broken = tmp_path / "broken.xlsx"
    workbook.save(broken)

    with pytest.raises(LayoutError):
        extract_workbook(broken, config)


def test_too_many_lessons_lost_fails(sheets: dict[str, list[Group]]) -> None:
    previous = previous_from(sheets)
    smaller = copy.deepcopy(sheets)
    for groups in smaller.values():
        for g in groups:
            g.lessons = g.lessons[:2]
    errors = validate(smaller, previous)
    assert any("Число занятий" in e for e in errors)


def test_group_without_lessons_fails(sheets: dict[str, list[Group]]) -> None:
    broken = copy.deepcopy(sheets)
    broken["ОФО"][0].lessons = []
    assert any("нет ни одного занятия" in e for e in validate(broken, None))


def test_diff_finds_changed_room(sheets: dict[str, list[Group]]) -> None:
    previous = previous_from(sheets)
    changed = copy.deepcopy(sheets)
    group = next(g for g in changed["ОФО"] if g.id == "ofo-1-bi-25")
    lesson = group.lessons[0]
    lesson.room = "9-99"

    diff = compare(previous, group_documents(changed, "new"), "new")
    assert diff["summary"] == {"added": 0, "removed": 0, "changed": 1}
    [change] = diff["groups"]["ofo-1-bi-25"]["changed"]
    assert change["changes"] == {"room": ["3-01", "9-99"]}


def test_diff_without_previous_is_empty(sheets: dict[str, list[Group]]) -> None:
    diff = compare(None, group_documents(sheets, "v1"), "v1")
    assert diff["from"] is None
    assert diff["groups"] == {}
