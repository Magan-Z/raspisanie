"""MergedSheet — удобный доступ к листу Excel с учётом объединённых ячеек.

В Excel у объединённой ячейки значение хранится только в левой верхней клетке,
а остальные клетки пустые. MergedSheet прячет эту особенность: для любой клетки
можно узнать её «область» (весь объединённый прямоугольник) и текст.
"""

from dataclasses import dataclass

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


@dataclass(frozen=True)
class Area:
    """Прямоугольник ячеек: строки top..bottom, столбцы left..right (включительно)."""
    top: int
    left: int
    bottom: int
    right: int

    @property
    def cols(self) -> range:
        return range(self.left, self.right + 1)

    @property
    def rows(self) -> range:
        return range(self.top, self.bottom + 1)

    @property
    def coord(self) -> str:
        """Адрес в стиле Excel: «AA8» или «AA8:AF9»."""
        start = f"{get_column_letter(self.left)}{self.top}"
        if self.top == self.bottom and self.left == self.right:
            return start
        return f"{start}:{get_column_letter(self.right)}{self.bottom}"


def clean_text(value: object) -> str:
    """Схлопывает переносы строк и повторные пробелы в один пробел, обрезает края."""
    if value is None:
        return ""
    return " ".join(str(value).split())


class MergedSheet:
    def __init__(self, ws: Worksheet) -> None:
        self.title = ws.title
        self.max_row = ws.max_row
        self.max_col = ws.max_column

        # Значения всех непустых клеток: (строка, столбец) → значение
        self._values: dict[tuple[int, int], object] = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    self._values[(cell.row, cell.column)] = cell.value

        # Для каждой клетки внутри объединения запоминаем её область
        self._areas: dict[tuple[int, int], Area] = {}
        for rng in ws.merged_cells.ranges:
            area = Area(rng.min_row, rng.min_col, rng.max_row, rng.max_col)
            for r in area.rows:
                for c in area.cols:
                    self._areas[(r, c)] = area

    def area(self, row: int, col: int) -> Area:
        """Область объединения, в которую входит клетка (или сама клетка)."""
        return self._areas.get((row, col)) or Area(row, col, row, col)

    def is_top_left(self, row: int, col: int) -> bool:
        a = self.area(row, col)
        return a.top == row and a.left == col

    def text(self, row: int, col: int) -> str:
        """Текст клетки: значение левой верхней клетки её области, без лишних пробелов."""
        a = self.area(row, col)
        return clean_text(self._values.get((a.top, a.left)))
